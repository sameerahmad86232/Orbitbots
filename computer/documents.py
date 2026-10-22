import base64,io,re
from fastapi import Depends,HTTPException
from pydantic import BaseModel,Field
class Extract(BaseModel):filename:str=Field(max_length=150);data:str=Field(max_length=14000000)
def install(s):
 @s.app.post('/documents/extract',dependencies=[Depends(s.reviewer)])
 def extract(body:Extract):
  try:raw=base64.b64decode(body.data,validate=True)
  except Exception:raise HTTPException(422,'Invalid base64 file')
  if len(raw)>10000000:raise HTTPException(413,'Document limit: 10 MB')
  name=body.filename.lower();parts=[];truncated=False
  try:
   if name.endswith('.pdf'):
    from pypdf import PdfReader
    doc=PdfReader(io.BytesIO(raw))
    if doc.is_encrypted:raise ValueError('Encrypted PDF')
    for page in doc.pages[:200]:
     parts.append(page.extract_text() or '')
     if sum(map(len,parts))>1000000:truncated=True;break
    truncated=truncated or len(doc.pages)>200
   elif name.endswith('.xlsx'):
    import zipfile
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
     if len(z.infolist())>5000 or sum(e.file_size for e in z.infolist())>50000000:raise ValueError('Expanded workbook too large')
    from openpyxl import load_workbook
    doc=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);cells=0
    try:
     for sheet in doc:
      parts.append('Worksheet: '+sheet.title)
      for row in sheet.iter_rows(values_only=True):
       cells+=len(row)
       if cells>100000:truncated=True;break
       parts.append('\t'.join('' if v is None else str(v) for v in row))
      if truncated:break
    finally:doc.close()
   else:parts=[raw.decode('utf-8')]
  except Exception:raise HTTPException(422,'Unsupported, encrypted or invalid document; no macros or OCR are executed')
  text='\n'.join(parts);truncated=truncated or len(text)>1000000
  return {'filename':body.filename,'text':text[:1000000],'truncated':truncated,'note':'Text/cached values only; no PDF OCR or formula execution'}
